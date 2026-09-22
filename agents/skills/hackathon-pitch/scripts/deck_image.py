#!/usr/bin/env python3
"""Drop a screenshot or diagram into a slide, inside an existing shape's box.

Requires python-pptx (`py -m pip install python-pptx`).

  py deck_image.py deck.pptx --slide 6 --shape 12 --image demo.png --out out.pptx

The target shape (find its id with `deck_text.py dump`) defines the frame: the
image is scaled to fit inside it, centred, aspect ratio preserved, and the
placeholder shape is removed. Use it for demo screenshots, architecture
diagrams, and charts rendered outside PowerPoint.

  --keep-shape   leave the placeholder in place (image goes on top)
  --cover        fill the box completely and crop the overflow instead of fitting
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is missing:  py -m pip install python-pptx")


def find_shape(slide, shape_id: int):
    for shape in slide.shapes:
        if shape.shape_id == int(shape_id):
            return shape
    return None


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--slide", required=True, type=int)
    ap.add_argument("--shape", required=True, type=int)
    ap.add_argument("--image", required=True, type=Path)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--keep-shape", action="store_true")
    ap.add_argument("--cover", action="store_true")
    a = ap.parse_args()

    target = a.pptx
    if a.out:
        a.out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(a.pptx, a.out)
        target = a.out

    prs = Presentation(str(target))
    slides = list(prs.slides)
    if not 1 <= a.slide <= len(slides):
        sys.exit(f"slide {a.slide} outside 1..{len(slides)}")
    slide = slides[a.slide - 1]

    box = find_shape(slide, a.shape)
    if box is None:
        sys.exit(f"no shape {a.shape} on slide {a.slide} — run deck_text.py dump")
    left, top, width, height = box.left, box.top, box.width, box.height

    pic = slide.shapes.add_picture(str(a.image), left, top)
    scale_w, scale_h = width / pic.width, height / pic.height
    scale = max(scale_w, scale_h) if a.cover else min(scale_w, scale_h)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = left + int((width - pic.width) / 2)
    pic.top = top + int((height - pic.height) / 2)

    if a.cover:
        # crop the overflow so the picture stays inside the original frame
        if pic.width > width:
            over = (pic.width - width) / pic.width / 2
            pic.crop_left = pic.crop_right = over
            pic.left, pic.width = left, width
        if pic.height > height:
            over = (pic.height - height) / pic.height / 2
            pic.crop_top = pic.crop_bottom = over
            pic.top, pic.height = top, height

    if not a.keep_shape:
        box._element.getparent().remove(box._element)

    prs.save(str(target))
    print(
        f"{target}  slide {a.slide}: {a.image.name} "
        f"{Emu(pic.width).cm:.1f}x{Emu(pic.height).cm:.1f}cm"
    )


if __name__ == "__main__":
    main()
