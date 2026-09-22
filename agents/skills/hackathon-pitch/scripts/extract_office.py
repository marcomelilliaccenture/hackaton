#!/usr/bin/env python3
"""Extract auditable content from an Office document (.docx / .pptx / .xlsx).

Produces a single JSON object on stdout containing everything the audit needs:
  - meta:     document type, counts (words, paragraphs, slides, sheets)
  - text:     full readable text, with structure markers, for anti-slop / readability / voice
  - segments: ordered list of {loc, text} blocks for precise location references
  - headings: list of {loc, level, text} where detectable
    - fonts/colors: included only with --include-visual for an explicitly requested visual audit

Default extraction is content-only. It does not inspect fonts, colors, fills, or other
visual properties. The --include-visual flag enables the style inventory only when the
user directly requested visual-brand or formatting review.

Usage:
    python3 extract_office.py <path-to-file> [--text-only] [--include-visual]

--text-only prints just the readable text (handy for a quick readability pass).

Dependencies (install once if missing):
    python3 -m pip install python-docx python-pptx openpyxl
The script reports a clear pip command if a needed library is absent.
"""
import json
import os
import re
import sys


def _fail(msg, code=1):
    print(json.dumps({"error": msg}), file=sys.stdout)
    sys.exit(code)


def _norm_hex(value):
    """Normalize a color value to an uppercase 6-digit #RRGGBB string, or None."""
    if value is None:
        return None
    s = str(value).strip().upper().lstrip("#")
    # openpyxl often gives 8-digit ARGB (e.g. FFA100FF); drop the alpha.
    if len(s) == 8:
        s = s[2:]
    if len(s) == 6 and re.fullmatch(r"[0-9A-F]{6}", s):
        return "#" + s
    return None


def _bump_color(colors, hex_value, context):
    if not hex_value:
        return
    entry = colors.setdefault(hex_value, {"count": 0, "contexts": []})
    entry["count"] += 1
    if context and context not in entry["contexts"]:
        entry["contexts"].append(context)


def _bump_font(fonts, name):
    if name:
        fonts[name] = fonts.get(name, 0) + 1


def _word_count(text):
    return len(re.findall(r"\b\w+\b", text, flags=re.UNICODE))


# --------------------------------------------------------------------------- DOCX
def extract_docx(path, include_visual=False):
    try:
        from docx import Document
    except ImportError:
        _fail("python-docx not installed. Run: python3 -m pip install python-docx")

    doc = Document(path)
    segments, headings, fonts, colors = [], [], {}, {}
    text_parts = []

    if include_visual:
        # Word usually sets the font at the style level, not per run.
        try:
            default_font = doc.styles["Normal"].font.name
            if default_font:
                fonts[default_font] = fonts.get(default_font, 0)
        except Exception:
            pass

    for i, para in enumerate(doc.paragraphs):
        ptext = para.text.strip()
        style_name = (para.style.name if para.style else "") or ""
        if include_visual:
            try:
                _bump_font(fonts, para.style.font.name)
            except Exception:
                pass
            for run in para.runs:
                _bump_font(fonts, run.font.name)
                try:
                    rgb = run.font.color.rgb if run.font.color and run.font.color.type is not None else None
                except Exception:
                    rgb = None
                _bump_color(colors, _norm_hex(rgb), "body run")
        if not ptext:
            continue
        loc = f"para {i + 1}"
        if style_name.lower().startswith("heading") or style_name.lower() == "title":
            level = "title" if style_name.lower() == "title" else style_name.split()[-1]
            headings.append({"loc": loc, "level": str(level), "text": ptext})
            text_parts.append(f"\n## [{style_name}] {ptext}\n")
        else:
            text_parts.append(ptext)
        segments.append({"loc": loc, "text": ptext})

    # Tables
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                loc = f"table {ti + 1} row {ri + 1}"
                segments.append({"loc": loc, "text": line})
                text_parts.append(line)

    full_text = "\n".join(text_parts)
    meta = {
        "type": "docx",
        "paragraphs": len([p for p in doc.paragraphs if p.text.strip()]),
        "tables": len(doc.tables),
        "words": _word_count(full_text),
    }
    return meta, full_text, segments, fonts, colors, headings


# --------------------------------------------------------------------------- PPTX
def extract_pptx(path, include_visual=False):
    try:
        from pptx import Presentation
        from pptx.util import Emu  # noqa: F401 (import validates the lib)
    except ImportError:
        _fail("python-pptx not installed. Run: python3 -m pip install python-pptx")

    prs = Presentation(path)
    segments, headings, fonts, colors = [], [], {}, {}
    text_parts = []

    for si, slide in enumerate(prs.slides):
        slide_no = si + 1
        text_parts.append(f"\n## [Slide {slide_no}]\n")
        for shape in slide.shapes:
            if include_visual:
                try:
                    if shape.fill.type is not None and shape.fill.type == 1:  # solid
                        _bump_color(colors, _norm_hex(shape.fill.fore_color.rgb), "shape fill")
                except Exception:
                    pass
            if not shape.has_text_frame:
                continue
            is_title = bool(getattr(shape, "is_placeholder", False) and
                            getattr(shape.placeholder_format, "idx", None) == 0)
            for pi, para in enumerate(shape.text_frame.paragraphs):
                runs_text = "".join(r.text for r in para.runs).strip()
                if include_visual:
                    for run in para.runs:
                        _bump_font(fonts, run.font.name)
                        try:
                            if run.font.color and run.font.color.type is not None:
                                _bump_color(colors, _norm_hex(run.font.color.rgb), "text run")
                        except Exception:
                            pass
                if not runs_text:
                    continue
                loc = f"slide {slide_no}"
                if is_title and pi == 0:
                    headings.append({"loc": loc, "level": "title", "text": runs_text})
                    text_parts.append(f"# {runs_text}")
                else:
                    text_parts.append(runs_text)
                segments.append({"loc": loc, "text": runs_text})

    full_text = "\n".join(text_parts)
    meta = {
        "type": "pptx",
        "slides": len(prs.slides),
        "words": _word_count(full_text),
    }
    return meta, full_text, segments, fonts, colors, headings


# --------------------------------------------------------------------------- XLSX
def extract_xlsx(path, include_visual=False):
    try:
        import openpyxl
    except ImportError:
        _fail("openpyxl not installed. Run: python3 -m pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    segments, headings, fonts, colors = [], [], {}, {}
    text_parts = []

    for ws in wb.worksheets:
        text_parts.append(f"\n## [Sheet: {ws.title}]\n")
        row_count = 0
        for row in ws.iter_rows():
            row_count += 1
            if row_count > 500:  # cap huge sheets — audit reads labels/headers, not every cell
                break
            values = []
            for cell in row:
                if cell.value is None:
                    continue
                if include_visual and cell.font:
                    _bump_font(fonts, cell.font.name)
                    try:
                        _bump_color(colors, _norm_hex(getattr(cell.font.color, "rgb", None)), "cell font")
                    except Exception:
                        pass
                if include_visual:
                    try:
                        if cell.fill and cell.fill.fgColor and cell.fill.patternType:
                            _bump_color(colors, _norm_hex(cell.fill.fgColor.rgb), "cell fill")
                    except Exception:
                        pass
                values.append(str(cell.value).strip())
            line = " | ".join(v for v in values if v)
            if line:
                loc = f"{ws.title}!row {cell.row}"
                segments.append({"loc": loc, "text": line})
                text_parts.append(line)

    full_text = "\n".join(text_parts)
    meta = {
        "type": "xlsx",
        "sheets": len(wb.worksheets),
        "words": _word_count(full_text),
    }
    return meta, full_text, segments, fonts, colors, headings


EXTRACTORS = {".docx": extract_docx, ".pptx": extract_pptx, ".xlsx": extract_xlsx}


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if not args:
        _fail("Usage: python3 extract_office.py <file.docx|.pptx|.xlsx> [--text-only] [--include-visual]")
    path = args[0]
    if not os.path.isfile(path):
        _fail(f"File not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext not in EXTRACTORS:
        _fail(f"Unsupported format '{ext}'. Supported: .docx, .pptx, .xlsx")

    include_visual = "--include-visual" in flags
    meta, full_text, segments, fonts, colors, headings = EXTRACTORS[ext](
        path, include_visual=include_visual
    )

    if "--text-only" in flags:
        print(full_text)
        return

    out = {
        "meta": {**meta, "file": os.path.basename(path)},
        "text": full_text,
        "segments": segments,
        "headings": headings,
    }
    if include_visual:
        out["fonts"] = dict(sorted(fonts.items(), key=lambda kv: -kv[1]))
        out["colors"] = dict(sorted(colors.items(), key=lambda kv: -kv[1]["count"]))
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
